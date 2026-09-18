"""CSV/Excel import with preview and validation."""
import csv
import io
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from sqlalchemy.orm import Session

from app.auth.jwt_handler import get_current_user
from app.database import get_db
from app.models.audit import AuditLog
from app.models.mine import Mine
from app.models.user import User, Subsidiary

router = APIRouter(prefix="/import", tags=["Import"])


def _rows_from_upload(content: bytes, filename: str) -> list[dict]:
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xls"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
            out = []
            for row in rows_iter:
                out.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
            return out
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Unable to read Excel file: {e}")
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return list(reader)


def _sanitize_cell(val):
    """Strip CSV/formula injection prefixes from cell values."""
    if isinstance(val, str):
        # Strip leading characters that could trigger formula injection in spreadsheets
        while val and val[0] in ('=', '+', '-', '@', '\t', '\r'):
            val = val[1:]
    return val


def _norm(row: dict) -> dict:
    return {str(k).strip().lower().replace(" ", "_"): _sanitize_cell(v) for k, v in row.items() if k}


@router.post("/mines/preview")
async def preview_mines(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if current_user.role not in ("admin", "analyst", "mine_operator"):
        raise HTTPException(status_code=403, detail="Not permitted")
    # File security validations
    fname = (file.filename or "upload.csv").lower()
    allowed_ext = (".csv", ".xlsx", ".xls")
    if not any(fname.endswith(ext) for ext in allowed_ext):
        raise HTTPException(status_code=400, detail=f"Unsupported file type. Allowed: {', '.join(allowed_ext)}")
    content = await file.read()
    max_size = 10 * 1024 * 1024  # 10 MB
    if len(content) > max_size:
        raise HTTPException(status_code=400, detail="File too large. Maximum size: 10 MB")
    raw = _rows_from_upload(content, file.filename or "upload.csv")
    valid, invalid, warnings = [], [], []
    codes = {m.code for m in db.query(Mine).all() if m.code}
    names = {m.name.lower() for m in db.query(Mine).all()}
    subs = {s.code.lower(): s.id for s in db.query(Subsidiary).all()}
    for i, row in enumerate(raw, start=2):
        r = _norm(row)
        name = str(r.get("name") or r.get("mine_name") or "").strip()
        code = str(r.get("code") or r.get("mine_code") or "").strip()
        errors = []
        if not name:
            errors.append("Name is required")
        lat = r.get("latitude")
        lon = r.get("longitude")
        try:
            if lat not in (None, ""):
                float(lat)
            if lon not in (None, ""):
                float(lon)
        except (TypeError, ValueError):
            errors.append("Invalid coordinates")
        if code and code in codes:
            errors.append("Duplicate mine code")
        if name.lower() in names:
            warnings.append(f"Row {i}: mine name already exists")
        item = {"row": i, "name": name, "code": code, "data": r, "errors": errors}
        if errors:
            invalid.append(item)
        else:
            valid.append(item)
    return {"valid": valid, "invalid": invalid, "warnings": warnings, "subsidiary_codes": list(subs.keys())}


@router.post("/mines/confirm")
async def confirm_mines(
    payload: dict,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if current_user.role != "admin":
        raise HTTPException(status_code=403, detail="Admin required to confirm import")
    rows = payload.get("rows") or []
    created = 0
    default_sub = db.query(Subsidiary).first()
    subs = {s.code.lower(): s.id for s in db.query(Subsidiary).all()}
    for item in rows:
        r = item.get("data") or {}
        name = str(r.get("name") or r.get("mine_name") or "").strip()
        if not name:
            continue
        sub_code = str(r.get("subsidiary") or r.get("subsidiary_code") or "").lower()
        mine = Mine(
            name=name,
            code=str(r.get("code") or f"IMP-{created + 1:03d}"),
            subsidiary_id=subs.get(sub_code) or (default_sub.id if default_sub else 1),
            location=str(r.get("location") or ""),
            district=str(r.get("district") or ""),
            state=str(r.get("state") or "Jharkhand"),
            latitude=float(r["latitude"]) if r.get("latitude") not in (None, "") else 23.79,
            longitude=float(r["longitude"]) if r.get("longitude") not in (None, "") else 86.43,
            mine_type=str(r.get("mine_type") or r.get("type") or "Opencast"),
            manager_name=str(r.get("manager") or r.get("manager_name") or ""),
            status=str(r.get("status") or "operational"),
            data_classification="demo_illustrative",
            source_name="User CSV/Excel import",
            retrieved_date=datetime.utcnow().date().isoformat(),
        )
        db.add(mine)
        created += 1
    db.add(AuditLog(
        user_id=current_user.id, user_name=current_user.name, user_role=current_user.role,
        action="imported", module="mines", entity="mine", details=f"Imported {created} mines", status="success",
    ))
    db.commit()
    return {"created": created}
